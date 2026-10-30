import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

import xlwt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas


class MarkItDownRuntimeTest(unittest.TestCase):
    def test_all_frozen_formats_convert_to_text(self):
        fixtures = {
            ".pdf": (make_pdf, "PDFCONVERTERTEXT"),
            ".docx": (make_docx, "DOCXCONVERTERTEXT"),
            ".pptx": (make_pptx, "PPTXCONVERTERTEXT"),
            ".xlsx": (make_xlsx, "XLSXCONVERTERTEXT"),
            ".xls": (make_xls, "XLSCONVERTERTEXT"),
        }
        with tempfile.TemporaryDirectory(prefix="lwc-document-import-test-") as directory:
            for extension, (writer, expected) in fixtures.items():
                with self.subTest(extension=extension):
                    source = Path(directory, "sample" + extension)
                    writer(source)
                    result = subprocess.run(
                        ["python", str(Path(__file__).with_name("convert.py")), str(source)],
                        check=False,
                        capture_output=True,
                        text=True,
                        timeout=45,
                    )
                    self.assertEqual(result.returncode, 0, f"{extension}: converter exited {result.returncode}")
                    self.assertIn(expected, result.stdout)

    def test_ocr_and_third_party_plugins_are_not_enabled(self):
        result = subprocess.run(
            ["python", "-m", "markitdown", "--list-plugins"],
            check=False,
            capture_output=True,
            text=True,
            timeout=10,
        )
        self.assertEqual(result.returncode, 0)
        self.assertIn("No 3rd-party plugins installed", result.stdout)

    def test_converter_process_network_guard_blocks_connections(self):
        result = subprocess.run(
            ["python", str(Path(__file__).with_name("convert.py")), "--network-guard-selftest"],
            check=False,
            capture_output=True,
            text=True,
            timeout=10,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("network blocked", result.stdout)


def make_pdf(path):
    writer = canvas.Canvas(str(path))
    writer.drawString(72, 720, "PDFCONVERTERTEXT")
    writer.save()


def make_docx(path):
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>'
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        '</Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body><w:p><w:r><w:t>DOCXCONVERTERTEXT</w:t></w:r></w:p>'
        '<w:sectPr/></w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", content_types)
        package.writestr("_rels/.rels", relationships)
        package.writestr("word/document.xml", document)


def make_pptx(path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    text_box.text_frame.text = "PPTXCONVERTERTEXT"
    deck.save(path)


def make_xlsx(path):
    workbook = Workbook()
    workbook.active["A1"] = "XLSXCONVERTERTEXT"
    workbook.save(path)


def make_xls(path):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet1")
    sheet.write(0, 0, "XLSCONVERTERTEXT")
    workbook.save(str(path))


if __name__ == "__main__":
    unittest.main()
